from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

ROOT = Path.cwd()
DATA = ROOT / 'data'
DRAFT = pd.read_csv(DATA / 'scoring_provisional_data.csv', encoding='utf-8-sig')
OFFICIAL = pd.read_csv(DATA / 'scoring_data.csv', encoding='utf-8-sig')

PILLARS = [
    ('Khách hàng', 'Customer_Score', ['C1','C2','C3']),
    ('Chiến lược', 'Strategy_Score', ['S1','S2','S3']),
    ('Công nghệ', 'Technology_Score', ['T1','T2','T3','T4']),
    ('Vận hành', 'Operations_Score', ['O1','O2','O3']),
    ('Văn hóa', 'Culture_Score', ['H1','H2','H3']),
    ('Dữ liệu', 'Data_Score', ['D1','D2','D3']),
]
CRIT_NAMES = {
'C1':'Bao phủ kênh và dịch vụ số', 'C2':'Mức độ khách hàng sử dụng kênh số', 'C3':'Hỗ trợ khách hàng trên kênh số',
'S1':'Chiến lược và lộ trình chuyển đổi số', 'S2':'Đầu tư/ngân sách cho chuyển đổi số', 'S3':'Hệ sinh thái và đối tác số',
'T1':'eKYC và sinh trắc học', 'T2':'Cho vay số/digital lending', 'T3':'AI/GenAI', 'T4':'Open Banking/API',
'O1':'Tỷ trọng giao dịch trên kênh số', 'O2':'Số hóa/tự động hóa quy trình', 'O3':'Vận hành dịch vụ số',
'H1':'Đào tạo và năng lực số', 'H2':'Đổi mới sáng tạo', 'H3':'Quản trị thay đổi và văn hóa số',
'D1':'Kiến trúc và tích hợp dữ liệu', 'D2':'Quản trị và an toàn dữ liệu', 'D3':'Phân tích dữ liệu và cá nhân hóa'
}
BANKS = ['Vietcombank','ACB','OCB','Agribank','VPBank']
SHORT = {'Vietcombank':'Vietcombank','ACB':'ACB','OCB':'OCB','Agribank':'Agribank','VPBank':'VPBank'}

NAVY='17324A'; TEAL='0F6B78'; SKY='DCEEF3'; LIGHT='F1F6F8'; BORDER='D7E2E9'; GRAY='526B7D'; WHITE='FFFFFF'; ORANGE='FFF4E5'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:'+m))
        if node is None:
            node = OxmlElement('w:'+m); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def format_table(table, header=True, font_size=8.5, first_col_bold=False):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    table.autofit = True
    for ri, row in enumerate(table.rows):
        if ri == 0 and header:
            set_repeat_table_header(row)
        for ci, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ri == 0 and header:
                shade(cell, TEAL)
            elif ri % 2 == 0:
                shade(cell, LIGHT)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                for run in p.runs:
                    run.font.name='Arial'; run.font.size=Pt(font_size)
                    run.font.color.rgb=RGBColor.from_string(WHITE if ri==0 and header else NAVY)
                    if ri==0 and header: run.bold=True
                    if ci==0 and first_col_bold: run.bold=True

def add_table(doc, headers, rows, font_size=8.5, first_col_bold=False):
    table=doc.add_table(rows=1, cols=len(headers))
    for cell, text in zip(table.rows[0].cells, headers): cell.text=str(text)
    for row in rows:
        cells=table.add_row().cells
        for cell, value in zip(cells,row): cell.text=str(value)
    format_table(table, font_size=font_size, first_col_bold=first_col_bold)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return table

def add_hyperlink(paragraph, text, url):
    part=paragraph.part
    rid=part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
    hyperlink=OxmlElement('w:hyperlink'); hyperlink.set(qn('r:id'),rid)
    run=OxmlElement('w:r'); rPr=OxmlElement('w:rPr')
    color=OxmlElement('w:color'); color.set(qn('w:val'),'176B87'); rPr.append(color)
    underline=OxmlElement('w:u'); underline.set(qn('w:val'),'single'); rPr.append(underline)
    run.append(rPr); t=OxmlElement('w:t'); t.text=text; run.append(t); hyperlink.append(run); paragraph._p.append(hyperlink)

def add_para(doc, text='', style=None, bold_lead=None):
    p=doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r=p.add_run(bold_lead); r.bold=True
        p.add_run(text[len(bold_lead):])
    else: p.add_run(text)
    return p

def fmt_score(v, digits=2):
    if pd.isna(v): return 'N/D'
    return f'{float(v):.{digits}f}'.replace('.', ',')

def cell_value(row, col):
    return row[col] if col in row.index else float('nan')

def page_field(paragraph):
    r=paragraph.add_run()
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)

# Document and typographic foundation
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(9.5); normal.font.color.rgb=RGBColor.from_string(NAVY)
normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.12
for sty, size, color in [('Title',25,NAVY),('Heading 1',16,NAVY),('Heading 2',12,TEAL),('Heading 3',10,TEAL)]:
    s=styles[sty]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)
    s.paragraph_format.space_before=Pt(12 if sty!='Title' else 2); s.paragraph_format.space_after=Pt(5)
header=sec.header.paragraphs[0]
header.text='BÁO CÁO NGHIÊN CỨU  |  CHUYỂN ĐỔI SỐ NGÂN HÀNG 2025'
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: r.font.name='Arial'; r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=footer.add_run('DTI ngân hàng 2025  •  Trang '); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(GRAY)
page_field(footer)

# Cover/title block
title=doc.add_paragraph(style='Title'); title.add_run('Báo cáo chi tiết xếp hạng chuyển đổi số ngân hàng')
subtitle=doc.add_paragraph(); subtitle.paragraph_format.space_after=Pt(14)
r=subtitle.add_run('Kỳ đánh giá 2025  |  Vietcombank, ACB, OCB, Agribank và VPBank')
r.font.name='Arial'; r.font.size=Pt(12); r.font.color.rgb=RGBColor.from_string(TEAL)

# Executive summary callout
call=doc.add_table(rows=1,cols=1); c=call.cell(0,0); shade(c,SKY); set_cell_margins(c,160,190,160,190)
p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
r=p.add_run('KẾT LUẬN ĐIỀU HÀNH\n'); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL); r.font.size=Pt(10)
p.add_run('Bảng tham khảo xếp đủ 5 ngân hàng có điểm từ 65,83 đến 71,74. Vietcombank đứng đầu bảng tham khảo; khoảng cách giữa ngân hàng hạng 1 và hạng 5 là 5,91 điểm. Tuy nhiên, đây chưa phải bảng chính thức: kết quả chính thức hiện chỉ xếp được Vietcombank, Agribank và OCB; ACB và VPBank chưa đạt ngưỡng bằng chứng đã duyệt.')
format_table(call,header=False,font_size=9.5)

add_para(doc,'Báo cáo giải thích cách hình thành điểm, mức độ đầy đủ dữ liệu, thứ hạng tham khảo cho đủ mẫu, kết quả chính thức hiện có và những bước cần hoàn tất trước khi dùng kết quả làm kết luận học thuật.')

# 1 Scope
h=doc.add_heading('1. Mục tiêu và phạm vi',level=1)
add_para(doc,'Mục tiêu là so sánh mức độ chuyển đổi số của năm ngân hàng theo một bộ chỉ số thích nghi dành cho lĩnh vực ngân hàng. Kỳ dữ liệu là năm 2025. Bộ chỉ số gồm 19 tiêu chí, chia thành sáu trụ cột; nguồn chính là báo cáo thường niên, báo cáo phát triển bền vững và thông tin công bố chính thức của ngân hàng.')
add_para(doc,'Đây là điểm DTI do nhóm nghiên cứu xây dựng, tham chiếu cấu trúc sáu trụ cột của Quyết định 2158/QĐ-BTTTT. Kết quả không phải điểm DBI chính thức của Bộ Thông tin và Truyền thông, và không đại diện cho toàn bộ hệ thống ngân hàng Việt Nam.')

# 2 method
h=doc.add_heading('2. Bộ tiêu chí và cách chấm',level=1)
rows=[]
for pillar, col, codes in PILLARS:
    rows.append((pillar, ', '.join(codes), '; '.join(f'{c} {CRIT_NAMES[c]}' for c in codes)))
add_table(doc,['Trụ cột','Mã','Tiêu chí'],rows,font_size=8.2,first_col_bold=True)
add_para(doc,'Định tính. Nhóm quy đổi mức độ triển khai thành 0, 30, 50, 70 hoặc 100 điểm. Mức điểm cao hơn đòi hỏi bằng chứng về độ phủ, tích hợp hoặc kết quả vận hành; rubric là quy ước chấm của nhóm, không phải thang điểm được quy định sẵn trong Quyết định 2158.')
add_para(doc,'Định lượng. C2 (mức độ khách hàng sử dụng kênh số) và O1 (tỷ trọng giao dịch số) chỉ nên dùng khi con số có cùng ý nghĩa với tiêu chí và phạm vi/mẫu số được nêu rõ. Không suy tỷ lệ từ các con số khác phạm vi. Trong bảng tham khảo, số liệu ứng viên chưa xác nhận mẫu số vẫn cần được rà soát trước khi kết luận chính thức.')
add_para(doc,'Tính điểm. Điểm trụ cột là trung bình các tiêu chí có điểm trong trụ cột đó. DTI là trung bình sáu điểm trụ cột, mỗi trụ cột chiếm 1/6. N/D không được đổi thành 0; Coverage cho biết có bao nhiêu tiêu chí được chấm và không cộng vào DTI.')
add_para(doc,'Điều kiện xếp hạng của dự án là có ít nhất 12/19 tiêu chí và có điểm ở cả sáu trụ cột. Điều kiện này là quy tắc phân tích của nhóm. Bảng tham khảo cho phép nhìn đủ mẫu dựa trên bằng chứng ứng viên; bảng chính thức chỉ tính dòng Approved đạt các trường nguồn bắt buộc.')

# 3 result
h=doc.add_heading('3. Kết quả xếp hạng tham khảo',level=1)
add_para(doc,'Bảng dưới đây dùng điểm đề xuất ghép với các dòng bằng chứng Candidate có nguồn và vị trí tra cứu. Nó giúp so sánh đủ năm ngân hàng trong cùng một khung. Vì trạng thái Candidate chưa đồng nghĩa với đã đối chiếu độc lập, thứ hạng này cần được gọi là tham khảo hoặc sơ bộ.')
rank_rows=[]
for _,r in DRAFT.sort_values('Rank').iterrows():
    rank_rows.append([int(r['Rank']),r['Bank'],fmt_score(r['DTI_Total_Score']),f"{int(r['Criteria_Available'])}/19",f"{fmt_score(r['Data_Coverage'],1)}%",f"{int(r['Pillars_Available'])}/6"])
add_table(doc,['Hạng','Ngân hàng','DTI','Tiêu chí có điểm','Coverage','Trụ cột'],rank_rows,font_size=8.5,first_col_bold=True)

pillar_rows=[]
for _,r in DRAFT.sort_values('Rank').iterrows():
    pillar_rows.append([r['Bank']]+[fmt_score(cell_value(r,col),1) for _,col,_ in PILLARS]+[fmt_score(r['DTI_Total_Score'])])
add_table(doc,['Ngân hàng','Khách hàng','Chiến lược','Công nghệ','Vận hành','Văn hóa','Dữ liệu','DTI'],pillar_rows,font_size=7.8,first_col_bold=True)
add_para(doc,'Diễn giải. Ba ngân hàng ở giữa (ACB, OCB, Agribank) có DTI nằm trong khoảng 68,28–68,61, tức cách nhau chỉ khoảng 0,33 điểm; thứ tự của nhóm này nhạy với việc rà soát từng nguồn và cách chấm các tiêu chí. Không nên mô tả các chênh lệch nhỏ này như khác biệt lớn về năng lực.')

# Criterion matrices split by pillar
h=doc.add_heading('4. Bảng điểm theo 19 tiêu chí',level=1)
add_para(doc,'Các ô N/D là chưa có điểm trong bảng tham khảo. Những giá trị số là điểm đề xuất hiện được lưu trong dự án; các điểm có thể được cập nhật khi nhóm đối chiếu nguồn, phạm vi và rubric.')
for pillar,col,codes in PILLARS:
    doc.add_heading(pillar,level=2)
    matrix=[]
    for code in codes:
        vals=[]
        for bank in BANKS:
            rr=DRAFT[DRAFT.Bank.eq(bank)].iloc[0]
            vals.append(fmt_score(rr[code],1))
        matrix.append([code,CRIT_NAMES[code]]+vals)
    add_table(doc,['Mã','Tiêu chí']+[SHORT[b] for b in BANKS],matrix,font_size=7.8,first_col_bold=True)

# 5 bank-by-bank discussion
h=doc.add_heading('5. Phân tích từng ngân hàng',level=1)
interpretations={
'Vietcombank':('Vietcombank có DTI tham khảo cao nhất (71,74/100), dữ liệu cho 18/19 tiêu chí và đủ sáu trụ cột. Điểm Khách hàng 85,00 cao nhất trong các trụ cột của ngân hàng; Vận hành đạt 77,13. Dữ liệu có điểm thấp nhất trong sáu trụ cột (63,33), cho thấy cần rà soát thêm bằng chứng về kiến trúc, quản trị và khai thác dữ liệu. C2 vẫn N/D nên mức sử dụng kênh số không được thể hiện đầy đủ.'),
'ACB':('ACB đạt 68,61 điểm tham khảo, 17/19 tiêu chí và đủ sáu trụ cột. Khách hàng đạt 85,00 và Dữ liệu đạt 70,00; Công nghệ là trụ cột thấp nhất ở mức 60,00, phần lớn do các điểm công nghệ có mức triển khai khác nhau. C2 và O1 đang N/D trong bảng tham khảo; cần bổ sung số liệu khách hàng sử dụng và tỷ trọng giao dịch với phạm vi rõ trước khi kết luận.'),
'OCB':('OCB đạt 68,39 điểm, có 18/19 tiêu chí. Khách hàng đạt 76,00 và Vận hành 72,67; Chiến lược và Dữ liệu cùng ở 63,33. H3 là tiêu chí N/D. C2 và O1 có tỷ lệ đề xuất trong bảng tham khảo, nhưng phải kiểm chứng định nghĩa, mẫu số và phạm vi đo lường trong nguồn gốc trước khi coi là điểm chính thức.'),
'Agribank':('Agribank đạt 68,28 điểm tham khảo và có điểm đề xuất cho cả 19 tiêu chí. Vận hành đạt 79,00, là trụ cột cao nhất; Khách hàng đạt 62,33, thấp nhất trong sáu trụ cột của ngân hàng. C2=47 và O1=97 trong bảng tham khảo vẫn cần làm rõ định nghĩa và phạm vi công bố. Việc có điểm ở mọi ô không đồng nghĩa mọi quan sát đã được duyệt chính thức.'),
'VPBank':('VPBank đạt 65,83 điểm tham khảo, có 16/19 tiêu chí và đủ sáu trụ cột trong bảng tham khảo. Văn hóa đạt 70,00; Chiến lược thấp nhất ở 56,67, trong đó S2 được đề xuất 30 điểm. C2, O1 và H3 còn N/D. Đây là ngân hàng có điểm tham khảo thấp nhất trong mẫu, nhưng mức chênh 5,91 điểm so với Vietcombank chưa tự thân chứng minh khác biệt có ý nghĩa thống kê.')}
for bank in ['Vietcombank','ACB','OCB','Agribank','VPBank']:
    doc.add_heading(bank,level=2)
    rr=DRAFT[DRAFT.Bank.eq(bank)].iloc[0]
    doc.add_paragraph(interpretations[bank])
    vals=[f'{name}: {fmt_score(rr[col])}' for name,col,_ in PILLARS]
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2)
    lead=p.add_run('Điểm trụ cột: '); lead.bold=True; lead.font.color.rgb=RGBColor.from_string(TEAL)
    p.add_run(' | '.join(vals)+f". Coverage: {int(rr['Criteria_Available'])}/19 ({fmt_score(rr['Data_Coverage'],1)}%).")

# 6 official
h=doc.add_heading('6. Kết quả chính thức theo bằng chứng đã duyệt',level=1)
add_para(doc,'Bảng chính thức chỉ sử dụng bằng chứng Approved đã đáp ứng điều kiện nguồn, vị trí tra cứu và kỳ dữ liệu. Tại lần cập nhật báo cáo, ba ngân hàng đủ ngưỡng; ACB và VPBank chưa đủ số tiêu chí đã duyệt. Vì vậy, hai ngân hàng này không có DTI tổng và không được gán hạng trong bảng chính thức.')
off_rows=[]
for _,r in OFFICIAL.sort_values('Rank',na_position='last').iterrows():
    rank='—' if pd.isna(r['Rank']) else int(r['Rank'])
    off_rows.append([rank,r['Bank'],fmt_score(r['DTI_Total_Score']),fmt_score(r['Criteria_Available'],0)+'/19',fmt_score(r['Data_Coverage'],1)+'%',fmt_score(r['Pillars_Available'],0)+'/6'])
add_table(doc,['Hạng','Ngân hàng','DTI chính thức','Tiêu chí','Coverage','Trụ cột'],off_rows,font_size=8.5,first_col_bold=True)
add_para(doc,'Không được trộn hai bảng thành một kết luận duy nhất. Có thể trình bày thứ hạng tham khảo đủ năm ngân hàng để so sánh, nhưng cần ghi rõ trạng thái ứng viên. Khi hoàn tất đối chiếu, chuyển từng dòng phù hợp sang Approved và chạy lại pipeline; kết quả chính thức sẽ được cập nhật từ các dòng đã duyệt.')

# 7 interpretation and limitations
h=doc.add_heading('7. Cách hiểu N/D và các giới hạn',level=1)
for text in [
'N/D có nghĩa là dự án chưa có quan sát đạt điều kiện chấm trong tập dữ liệu hiện tại. N/D không phải 0 điểm và không chứng minh ngân hàng không có hoạt động đó.',
'Bảng tham khảo sử dụng điểm đề xuất và Candidate; tự động trích văn bản có thể bỏ sót ngữ cảnh, nhầm trang, hoặc ghép một đoạn chưa đủ mạnh với tiêu chí. Cần đọc báo cáo gốc trước khi xác nhận.',
'Các tiêu chí C2 và O1 cần đặc biệt thận trọng vì tỷ lệ chỉ so sánh được khi khái niệm, mẫu số, phạm vi khách hàng/giao dịch và kỳ đo giống nhau.',
'Điểm trung bình theo trụ cột có thể dựa trên số tiêu chí khác nhau giữa các ngân hàng. Vì vậy Coverage phải được đọc kèm DTI; điểm tổng cao hơn không đồng nghĩa mọi tiêu chí đều tốt hơn.',
'Mẫu chỉ có năm ngân hàng được chọn theo yêu cầu đề tài. Không suy rộng thứ hạng sang toàn ngành; bộ điểm là mô tả theo tài liệu công khai và kỳ 2025.'
]:
    doc.add_paragraph(text,style='List Bullet')

# 8 next steps
h=doc.add_heading('8. Việc cần làm để chốt kết quả',level=1)
steps=[
'Đối chiếu từng đoạn Candidate của ACB và VPBank với báo cáo thường niên 2025 bản gốc, kiểm tra trang và phạm vi; chỉ chuyển Approved khi đoạn trích thực sự chứng minh tiêu chí.',
'Rà soát các dòng C2/O1 của OCB và Agribank để xác nhận định nghĩa tỷ lệ, phạm vi đo và mẫu số. Nếu không rõ, giữ N/D trong kết quả chính thức.',
'Đối chiếu toàn bộ điểm định tính với rubric 0/30/50/70/100, đặc biệt các điểm S2 và các tiêu chí có mức triển khai chỉ ở một sản phẩm/đơn vị.',
'Chạy lại pipeline, kiểm tra số tiêu chí và sáu trụ cột của từng ngân hàng, rồi cập nhật bảng chính thức, bảng tham khảo và báo cáo cùng một phiên bản dữ liệu.',
'Khi trình bày, ghi rõ ngày chốt dữ liệu, kỳ đánh giá 2025, quy tắc N/D và khác biệt giữa kết quả tham khảo với kết quả chính thức.'
]
for i,s in enumerate(steps,1):
    p=doc.add_paragraph(style='List Number'); p.add_run(s)

# Sources from the project record and official annual reports
h=doc.add_heading('9. Nguồn dữ liệu chính',level=1)
add_para(doc,'Bằng chứng điểm được lưu cùng URL, trang, trích đoạn và trạng thái trong data/evidence_data.csv của dự án. Các nguồn báo cáo chính thức được liên kết dưới đây:')
sources=[
('Vietcombank','Báo cáo Phát triển bền vững 2025','https://www.vietcombank.com.vn/-/media/Project/VCB-Sites/VCB/Nha-Dau-tu/Bao-cao-PTBV/VN_VCB_ESG-Report-2025.pdf?ts=20260428074438'),
('VPBank','Báo cáo thường niên 2025','https://www.vpbank.com.vn/-/media/vpbank-latest/5nha-dau-tu/bao-cao-thuong-nien/2025/bctc-2025-0605.pdf'),
('ACB','Báo cáo thường niên 2025','https://acb.com.vn/nha-dau-tu/bao-cao-thuong-nien-2025'),
('OCB','Báo cáo thường niên 2025','https://webocb-api.ocb.com.vn/Resources/Files/20260408175516_20260408-ocb-cbtt-bao-cao-thuong-nien-nam-2025.pdf'),
('Agribank','Báo cáo thường niên 2025','https://www.agribank.com.vn/wcm/connect/ef81c291-19b0-4cc0-8f30-ff348c7d5f91/25_04_2026_VIE.pdf?CACHEID=ROOTWORKSPACE-ef81c291-19b0-4cc0-8f30-ff348c7d5f91-pThNs3K&CONVERT_TO=url&MOD=AJPERES'),
]
for bank,label,url in sources:
    p=doc.add_paragraph(style='List Bullet'); p.add_run(bank+' — '+label+': '); add_hyperlink(p,'mở nguồn chính thức',url)

# Final note
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10)
r=p.add_run('Tệp kết quả tham chiếu: '); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
p.add_run('data/scoring_provisional_data.csv (tham khảo), data/scoring_data.csv (chính thức), FINAL_REPORT.md (tóm tắt).')

out=ROOT/'Bao_cao_chi_tiet_xep_hang_CDS_ngan_hang_2025.docx'
doc.save(out)
print(out)
